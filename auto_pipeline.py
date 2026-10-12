import os
import re
import json
import requests
from bs4 import BeautifulSoup
from urllib.parse import urljoin, urlparse
import openpyxl
import pdfplumber
try:
    import docx
except ImportError:
    docx = None

DAYS = ["Pazartesi", "Salı", "Çarşamba", "Perşembe", "Cuma"]

DEPT_NAMES = {
    "pc": "Bilgisayar Mühendisliği",
    "eem": "Elektrik-Elektronik Mühendisliği",
    "makine": "Makine Mühendisliği",
    "mekatronik": "Mekatronik Mühendisliği",
    "mekatroniking": "Mekatronik Mühendisliği (İngilizce)",
    "endustri": "Endüstri Mühendisliği",
    "insaat": "İnşaat Mühendisliği",
    "kimyamuh": "Kimya Mühendisliği",
    "kimyamüh": "Kimya Mühendisliği",
    "kimya": "Kimya",
    "biyomuhendis": "Biyomühendislik",
    "gida_muh": "Gıda Mühendisliği",
    "yapayzeka": "Yapay Zeka ve Veri Mühendisliği",
    "veri": "Veri Bilimi ve Analitiği",
    "matematik": "Matematik",
    "fizik": "Fizik",
    "isletme": "İşletme"
}

DOWNLOAD_DIR = "kaynak_dosyalar"
os.makedirs(DOWNLOAD_DIR, exist_ok=True)

# -------------------------------------------------------------
# 1. MODÜL: WEB'DEN DERS PROGRAMI LİNKLERİNİ BUL VE İNDİR
# -------------------------------------------------------------
def download_file(url, target_path):
    try:
        headers = {"User-Agent": "Mozilla/5.0 (Windows NT 10.0; Win64; x64)"}
        resp = requests.get(url, headers=headers, timeout=20)
        if resp.status_code == 200:
            with open(target_path, "wb") as f:
                f.write(resp.content)
            print(f"[+] İndirildi: {os.path.basename(target_path)}")
            return True
    except Exception as e:
        print(f"[-] İndirme hatası ({url}): {e}")
    return False

def scrape_btu_programs(target_urls):
    """Belirtilen sayfalardan PDF, Excel ve Word linklerini toplar."""
    print("=== Web Sayfalarından Ders Programları Taranıyor ===")
    headers = {"User-Agent": "Mozilla/5.0 (Windows NT 10.0; Win64; x64)"}
    valid_exts = [".pdf", ".xlsx", ".xls", ".docx"]

    for page_url in target_urls:
        try:
            resp = requests.get(page_url, headers=headers, timeout=15)
            if resp.status_code != 200:
                continue
            soup = BeautifulSoup(resp.text, "html.parser")
            for a in soup.find_all("a", href=True):
                href = a["href"].strip()
                full_url = urljoin(page_url, href)
                path = urlparse(full_url).path
                ext = os.path.splitext(path)[1].lower()

                # Ders programı anahtar kelimeleri
                link_text = (a.get_text() + " " + href).lower()
                is_schedule = any(k in link_text for k in ["program", "ders", "timetable", "guz", "bahar"])

                if ext in valid_exts and is_schedule:
                    file_name = os.path.basename(path)
                    if not file_name:
                        continue
                    save_path = os.path.join(DOWNLOAD_DIR, file_name)
                    if not os.path.exists(save_path):
                        download_file(full_url, save_path)
        except Exception as e:
            print(f"[-] Tarama hatası ({page_url}): {e}")

# -------------------------------------------------------------
# 2. MODÜL: HÜCRE PARSING & DÜZENLEME
# -------------------------------------------------------------
def clean_whitespace(t):
    if not t: return ""
    return re.sub(r"\s+", " ", str(t)).strip()

def detect_department_from_filename(filename):
    base = os.path.splitext(filename)[0].lower().replace("-", "_").replace(" ", "_")
    m = re.match(r"^([a-z_çğıöşü]+?)(?:_?(\d))?$", base)
    dept_key = base
    grade = None
    if m:
        dept_key = m.group(1).rstrip("_")
        if m.group(2):
            grade = int(m.group(2))
    dept_name = DEPT_NAMES.get(dept_key) or dept_key.replace("_", " ").title()
    return dept_name, grade

def extract_grade_from_text(text):
    if not text: return None
    t = text.upper()
    for g in [1, 2, 3, 4]:
        if f"{g}. SINIF" in t or f"{g}.SINIF" in t:
            return g
    return None

def extract_hour(time_text):
    if not time_text: return None
    m = re.search(r"(\d{1,2}):\d{2}", str(time_text))
    return int(m.group(1)) if m else None

def parse_course_cell(cell_text):
    if not cell_text: return []
    text = clean_whitespace(cell_text)
    if len(text) < 4: return []

    parts = [p.strip() for p in text.split("■") if p.strip()] if "■" in text else [text]
    results = []

    for item in parts:
        code_match = re.search(r"\b([A-ZÇĞİÖŞÜ]{2,5}\s*\d{3,4})\b", item)
        code = code_match.group(1).replace(" ", "") if code_match else ""

        sec_match = re.search(r"(?:Şube|Sube)\s*(\d+)", item, re.IGNORECASE)
        section = f"Şube {sec_match.group(1)}" if sec_match else "Şube 1"

        room_match = re.search(r"((?:MS|YB)\s*/[^,]+|[A-Z]/\d+[^,]*|[\w\s]+Laboratuvar[\w\s\.]*|[\w\s]+Amfi[\w\s]*)", item, re.IGNORECASE)
        classroom = room_match.group(1).strip() if room_match else ""

        rem = item
        if code_match:
            rem = item[item.find(code_match.group(0)) + len(code_match.group(0)):].strip()
        name_part = rem.split(",")[0].strip()

        type_match = re.search(r"\s+([TL])$", name_part)
        type_suffix = ""
        if type_match:
            type_suffix = " (Teori)" if type_match.group(1) == "T" else " (Lab)"
            name_part = name_part[:type_match.start()].strip()

        course_name = f"{name_part}{type_suffix}" if name_part else (code or item)

        instructor = ""
        comma_parts = [p.strip() for p in item.split(",") if p.strip()]
        for p in comma_parts[1:]:
            if not re.search(r"(Şube|Sube|BLOK|LAB|MS\s*/|YB\s*/|PROGRAMI|Mühendis)", p, re.IGNORECASE):
                if any(c.isupper() for c in p) and len(p) > 3:
                    instructor = p
                    break

        if code or len(course_name) > 3:
            results.append({
                "raw": item,
                "code": code,
                "name": course_name,
                "section": section,
                "instructor": instructor,
                "classroom": classroom
            })
    return results

def process_table_matrix(matrix, default_dept, default_grade, page_num=1):
    day_indices = {}
    header_row_idx = -1
    detected_grade = default_grade

    for r_idx, row in enumerate(matrix):
        cleaned = [clean_whitespace(c) for c in row]
        row_txt = " ".join(cleaned)
        g = extract_grade_from_text(row_txt)
        if g: detected_grade = g

        found = {}
        for i, c in enumerate(cleaned):
            for d in DAYS:
                if d.lower() in c.lower():
                    found[d] = i
        if len(found) >= 3:
            day_indices = found
            header_row_idx = r_idx
            break

    if header_row_idx == -1:
        return None, []

    if not detected_grade:
        detected_grade = page_num if page_num <= 4 else 1

    program_key = f"{default_dept} - {detected_grade}. Sınıf"
    courses = []

    for row in matrix[header_row_idx + 1:]:
        if not row: continue
        time_cell = ""
        for cell in row[:2]:
            if cell and re.search(r"\d{1,2}:\d{2}", str(cell)):
                time_cell = str(cell)
                break
        hour = extract_hour(time_cell)
        if hour is None: continue

        for day, col_idx in day_indices.items():
            if col_idx < len(row):
                val = row[col_idx]
                if val:
                    for c in parse_course_cell(val):
                        courses.append({
                            **c,
                            "department": default_dept,
                            "grade": detected_grade,
                            "program": program_key,
                            "day": day,
                            "hour": hour,
                            "time": f"{hour:02d}:00 - {hour:02d}:45"
                        })
    return program_key, courses

# -------------------------------------------------------------
# 3. MODÜL: FORMAT AYRIŞTIRICILARI (PDF, EXCEL, DOCX)
# -------------------------------------------------------------
def parse_docx(fpath):
    if not docx: return []
    doc = docx.Document(fpath)
    matrices = []
    for table in doc.tables:
        matrix = []
        for row in table.rows:
            matrix.append([cell.text.strip() for cell in row.cells])
        matrices.append(matrix)
    return matrices

def parse_excel(fpath):
    matrices = []
    wb = openpyxl.load_workbook(fpath, data_only=True)
    for s_name in wb.sheetnames:
        sheet = wb[s_name]
        table = list(sheet.iter_rows(values_only=True))
        if table and len(table) >= 2:
            matrices.append((s_name, table))
    return matrices

def parse_pdf(fpath):
    matrices = []
    with pdfplumber.open(fpath) as pdf:
        for p_idx, page in enumerate(pdf.pages, start=1):
            tables = page.extract_tables({
                "vertical_strategy": "lines",
                "horizontal_strategy": "lines",
                "snap_tolerance": 5,
                "join_tolerance": 5
            })
            if not tables:
                tables = page.extract_tables()
            for t in tables:
                if t and len(t) >= 2:
                    matrices.append((p_idx, t))
    return matrices

# -------------------------------------------------------------
# 4. MODÜL: DERSLİK DOLULUK VERİTABANINI ÜRET (data.js)
# -------------------------------------------------------------
def build_classrooms_data(all_programs):
    """Tüm ders saatlerinden boş/dolu derslik matrisini oluşturur."""
    classrooms = {}
    for prog in all_programs.values():
        for c in prog["courses"]:
            room = c.get("classroom")
            if not room or len(room) < 3:
                continue
            if room not in classrooms:
                classrooms[room] = []
            
            # Mükerrer kayıt önleme
            exists = any(
                slot["day"] == c["day"] and slot["hour"] == c["hour"]
                for slot in classrooms[room]
            )
            if not exists:
                classrooms[room].append({
                    "day": c["day"],
                    "hour": c["hour"],
                    "start": f"{c['hour']:02d}:00",
                    "end": f"{c['hour']:02d}:45",
                    "status": "DOLU"
                })
    return classrooms

# -------------------------------------------------------------
# ANA ÇALIŞTIRICI
# -------------------------------------------------------------
def main():
    # 1. Örnek olarak taranacak okul duyuru sayfaları
    # İstediğin fakülte/bölüm duyuru linklerini buraya ekleyebilirsin
    BTU_PAGES = [
        "https://bm.btu.edu.tr/tr/duyurular",
        "https://eem.btu.edu.tr/tr/duyurular",
        "https://endustri.btu.edu.tr/tr/duyurular"
    ]
    # scrape_btu_programs(BTU_PAGES)  # İsteğe bağlı olarak web taramayı tetikler

    all_programs = {}
    total_slots = 0

    print("\n=== Klasördeki Tüm Dosyalar (PDF, Excel, Word) Ayrıştırılıyor ===")
    files = sorted(os.listdir(DOWNLOAD_DIR))

    for fname in files:
        fpath = os.path.join(DOWNLOAD_DIR, fname)
        if not os.path.isfile(fpath): continue
        ext = os.path.splitext(fname)[1].lower()
        dept_name, file_grade = detect_department_from_filename(fname)

        # PDF
        if ext == ".pdf":
            try:
                for p_idx, tbl in parse_pdf(fpath):
                    pkey, courses = process_table_matrix(tbl, dept_name, file_grade, page_num=p_idx)
                    if pkey and courses:
                        if pkey not in all_programs:
                            all_programs[pkey] = {"title": pkey, "department": dept_name, "grade": courses[0]["grade"], "file": fname, "courses": []}
                        all_programs[pkey]["courses"].extend(courses)
                        total_slots += len(courses)
            except Exception as e:
                print(f"[-] PDF Hatası ({fname}): {e}")

        # EXCEL
        elif ext in [".xlsx", ".xls"]:
            try:
                for s_name, tbl in parse_excel(fpath):
                    s_grade = extract_grade_from_text(s_name) or file_grade
                    pkey, courses = process_table_matrix(tbl, dept_name, s_grade)
                    if pkey and courses:
                        if pkey not in all_programs:
                            all_programs[pkey] = {"title": pkey, "department": dept_name, "grade": courses[0]["grade"], "file": fname, "courses": []}
                        all_programs[pkey]["courses"].extend(courses)
                        total_slots += len(courses)
            except Exception as e:
                print(f"[-] Excel Hatası ({fname}): {e}")

        # WORD (.docx)
        elif ext == ".docx":
            try:
                for idx, tbl in enumerate(parse_docx(fpath), start=1):
                    pkey, courses = process_table_matrix(tbl, dept_name, file_grade, page_num=idx)
                    if pkey and courses:
                        if pkey not in all_programs:
                            all_programs[pkey] = {"title": pkey, "department": dept_name, "grade": courses[0]["grade"], "file": fname, "courses": []}
                        all_programs[pkey]["courses"].extend(courses)
                        total_slots += len(courses)
            except Exception as e:
                print(f"[-] Word Hatası ({fname}): {e}")

    # 1. courses_data.js çıktısı
    with open("courses_data.js", "w", encoding="utf-8") as f:
        f.write("const ALL_PROGRAMS_DATA = " + json.dumps(all_programs, ensure_ascii=False, indent=2) + ";")

    # 2. data.js (Derslik Doluluk) çıktısı
    classrooms_data = build_classrooms_data(all_programs)
    with open("data.js", "w", encoding="utf-8") as f:
        f.write("const CLASSROOMS_DATA = " + json.dumps(classrooms_data, ensure_ascii=False, indent=2) + ";")

    # 3. Eklenti (extension) klasörünü de senkronize et
    if os.path.isdir("extension"):
        try:
            import shutil
            shutil.copyfile("courses_data.js", os.path.join("extension", "courses_data.js"))
            shutil.copyfile("data.js", os.path.join("extension", "data.js"))
            if os.path.exists("exam_schedules.js"):
                shutil.copyfile("exam_schedules.js", os.path.join("extension", "exam_schedules.js"))
        except Exception as e:
            print(f"[-] Extension senkronizasyon uyarisi: {e}")

    # 4. Sınav Takvimlerini Tara ve Güncelle
    try:
        from scrape_exams import main as scrape_exams_main
        print("\n=== Sınav Takvimleri Taranıyor ===")
        scrape_exams_main()
        if os.path.isdir("extension") and os.path.exists("exam_schedules.js"):
            import shutil
            shutil.copyfile("exam_schedules.js", os.path.join("extension", "exam_schedules.js"))
    except Exception as e:
        print(f"[-] Sınav takvimi tarama uyarisi: {e}")

    print("\n[✓] İŞLEM TAMAMLANDI!")
    print(f"Toplam Ayrıştırılan Bölüm Programı: {len(all_programs)}")
    print(f"Toplam Ders Saati Slotu: {total_slots}")
    print(f"Oluşturulan/Güncellenen Dosyalar: courses_data.js, data.js & exam_schedules.js")

if __name__ == "__main__":
    main()