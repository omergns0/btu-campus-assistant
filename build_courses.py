import os
import re
import json
import openpyxl
import pdfplumber

DAYS = ["Pazartesi", "Salı", "Çarşamba", "Perşembe", "Cuma"]

DEPT_NAMES = {
    "pc": "Bilgisayar Mühendisliği",
    "eem": "Elektrik-Elektronik Mühendisliği",
    "makine": "Makine Mühendisliği",
    "mekatronik": "Mekatronik Mühendisliği",
    "mekatroniking": "Mekatronik Mühendisliği (İngilizce)",
    "endustri": "Endüstri Mühendisliği",
    "insaat": "İnşaat Mühendisliği",
    "kimyamuh": "Kimya Mühendisliği",
    "kimyamüh": "Kimya Mühendisliği",
    "kimya": "Kimya",
    "biyomuhendis": "Biyomühendislik",
    "cevre_muhendis": "Çevre Mühendisliği",
    "çevre_mühendis": "Çevre Mühendisliği",
    "gida_muh": "Gıda Mühendisliği",
    "gemiinsaati": "Gemi İnşaatı ve Gemi Makineleri Mühendisliği",
    "meta_malzeme": "Metalurji ve Malzeme Mühendisliği",
    "polimer": "Polimer Malzeme Mühendisliği",
    "yapayzeka": "Yapay Zeka ve Veri Mühendisliği",
    "veri": "Veri Bilimi ve Analitiği",
    "matematik": "Matematik",
    "fizik": "Fizik",
    "isletme": "İşletme",
    "psikoloji": "Psikoloji",
    "sosyoloji": "Sosyoloji",
    "uluslariliski": "Uluslararası İlişkiler",
    "uluslariliski_ing": "Uluslararası İlişkiler (İngilizce)",
    "utl": "Uluslararası Ticaret ve Lojistik",
    "mutercim_ve_tercumanlik": "Mütercim ve Tercümanlık",
    "orman_muhendisligi": "Orman Mühendisliği",
    "orman_mühendisliği": "Orman Mühendisliği",
    "orman_endustri": "Orman Endüstri Mühendisliği",
    "pm": "Peyzaj Mimarlığı",
    "denizcilik_isletmeleri": "Denizcilik İşletmeleri Yönetimi"
}

def clean_whitespace(t):
    if not t: return ""
    return re.sub(r"\s+", " ", str(t)).strip()

def detect_department_from_filename(filename):
    base = os.path.splitext(filename)[0].lower().replace("-", "_").replace(" ", "_")
    m = re.match(r"^([a-z_çğıöşü]+?)(?:_?(\d))?$", base)
    dept_key = base
    grade = None
    if m:
        dept_key = m.group(1).rstrip("_")
        if m.group(2):
            grade = int(m.group(2))
    
    dept_name = DEPT_NAMES.get(dept_key) or dept_key.replace("_", " ").title()
    return dept_name, grade

def extract_grade_from_text(text):
    if not text: return None
    t = text.upper()
    if "1. SINIF" in t or "1.SINIF" in t: return 1
    if "2. SINIF" in t or "2.SINIF" in t: return 2
    if "3. SINIF" in t or "3.SINIF" in t: return 3
    if "4. SINIF" in t or "4.SINIF" in t: return 4
    return None

def parse_course_cell(cell_text):
    if not cell_text: return []
    text = clean_whitespace(cell_text)
    if len(text) < 4: return []

    parts = []
    if "■" in text:
        parts = [p.strip() for p in text.split("■") if p.strip()]
    else:
        # Kod deseni yakalama
        matches = list(re.finditer(r"\b([A-ZÇĞİÖŞÜ]{2,5}\s*\d{3,4})\b", text))
        if len(matches) > 1:
            for idx, m in enumerate(matches):
                start = m.start()
                end = matches[idx + 1].start() if idx + 1 < len(matches) else len(text)
                parts.append(text[start:end].strip())
        else:
            parts = [text]

    results = []
    for item in parts:
        code_match = re.search(r"\b([A-ZÇĞİÖŞÜ]{2,5}\s*\d{3,4})\b", item)
        code = code_match.group(1).replace(" ", "") if code_match else ""

        # Şube tespiti
        sec_match = re.search(r"(?:Şube|Sube)\s*(\d+)", item, re.IGNORECASE)
        section = f"Şube {sec_match.group(1)}" if sec_match else "Şube 1"

        # Derslik tespiti
        room_match = re.search(r"((?:MS|YB)\s*/[^,]+|[A-Z]/\d+[^,]*|[\w\s]+Laboratuvar[\w\s\.]*|[\w\s]+Amfi[\w\s]*)", item, re.IGNORECASE)
        classroom = room_match.group(1).strip() if room_match else ""

        # Ders adı tespiti
        rem = item
        if code_match:
            rem = item[item.find(code_match.group(0)) + len(code_match.group(0)):].strip()
        
        name_part = rem.split(",")[0].strip()

        # T / L tespiti
        type_suffix = ""
        type_match = re.search(r"\s+([TL])$", name_part)
        if type_match:
            t_type = type_match.group(1)
            type_suffix = " (Teori)" if t_type == "T" else " (Lab)"
            name_part = name_part[:type_match.start()].strip()

        course_name = f"{name_part}{type_suffix}" if name_part else (code or item)

        # Hoca tespiti
        instructor = ""
        comma_parts = [p.strip() for p in item.split(",") if p.strip()]
        for p in comma_parts[1:]:
            if not re.search(r"(Şube|Sube|BLOK|LAB|MS\s*/|YB\s*/|PROGRAMI|Mühendis)", p, re.IGNORECASE):
                if any(c.isupper() for c in p) and len(p) > 3:
                    instructor = p
                    break

        if code or len(course_name) > 3:
            results.append({
                "raw": item,
                "code": code,
                "name": course_name,
                "section": section,
                "instructor": instructor,
                "classroom": classroom
            })

    return results

def extract_hour(time_text):
    if not time_text: return None
    m = re.search(r"(\d{1,2}):\d{2}", str(time_text))
    return int(m.group(1)) if m else None

def process_table(table, default_dept, default_grade, page_num=1):
    day_indices = {}
    header_row_idx = -1
    detected_grade = default_grade

    for r_idx, row in enumerate(table):
        cleaned = [clean_whitespace(c) for c in row]
        row_txt = " ".join(cleaned)
        
        g = extract_grade_from_text(row_txt)
        if g: detected_grade = g

        found = {}
        for i, c in enumerate(cleaned):
            for d in DAYS:
                if d.lower() in c.lower():
                    found[d] = i
        
        if len(found) >= 3:
            day_indices = found
            header_row_idx = r_idx
            break

    if header_row_idx == -1:
        return None, []

    if not detected_grade:
        detected_grade = page_num if page_num <= 4 else 1

    program_key = f"{default_dept} - {detected_grade}. Sınıf"
    courses = []

    # Saat satırlarını tara
    for row in table[header_row_idx + 1:]:
        if not row: continue
        
        # Saat hücresi ilk iki sütundan birindedir
        time_cell = ""
        for cell in row[:2]:
            if cell and re.search(r"\d{1,2}:\d{2}", str(cell)):
                time_cell = str(cell)
                break
        
        hour = extract_hour(time_cell)
        if hour is None: continue

        for day, col_idx in day_indices.items():
            if col_idx < len(row):
                cell_val = row[col_idx]
                if cell_val:
                    parsed = parse_course_cell(cell_val)
                    for c in parsed:
                        courses.append({
                            **c,
                            "department": default_dept,
                            "grade": detected_grade,
                            "program": program_key,
                            "day": day,
                            "hour": hour,
                            "time": f"{hour:02d}:00 - {hour:02d}:45"
                        })

    return program_key, courses

def main():
    source_dir = "kaynak_dosyalar" if os.path.exists("kaynak_dosyalar") else "."
    all_programs = {}
    total_slots = 0

    print("=== BTÜ Ders Programları Lineer Cebir & Cuma Günü Düzeltmesiyle Taranıyor ===")
    files = sorted(os.listdir(source_dir))
    for fname in files:
        fpath = os.path.join(source_dir, fname)
        if not os.path.isfile(fpath): continue

        ext = os.path.splitext(fname)[1].lower()
        if ext not in [".pdf", ".xlsx", ".xls"]: continue

        dept_name, file_grade = detect_department_from_filename(fname)

        if ext == ".pdf":
            try:
                with pdfplumber.open(fpath) as pdf:
                    for p_idx, page in enumerate(pdf.pages, start=1):
                        # Tablo çıkarma toleranslarını hassaslaştırıyoruz (birleşik hücreleri kaçırmaz)
                        tables = page.extract_tables({
                            "vertical_strategy": "lines",
                            "horizontal_strategy": "lines",
                            "snap_tolerance": 5,
                            "join_tolerance": 5
                        })
                        if not tables:
                            tables = page.extract_tables() # Yedek strateji

                        for t in tables:
                            if not t or len(t) < 2: continue
                            prog_key, courses = process_table(t, dept_name, file_grade, page_num=p_idx)
                            if prog_key and courses:
                                if prog_key not in all_programs:
                                    all_programs[prog_key] = {
                                        "title": prog_key,
                                        "department": dept_name,
                                        "grade": courses[0]["grade"],
                                        "file": fname,
                                        "courses": []
                                    }
                                all_programs[prog_key]["courses"].extend(courses)
                                total_slots += len(courses)
            except Exception as e:
                print(f"  [-] PDF Hatası ({fname}): {e}")

        elif ext in [".xlsx", ".xls"]:
            try:
                wb = openpyxl.load_workbook(fpath, data_only=True)
                for s_idx, sheet_name in enumerate(wb.sheetnames, start=1):
                    sheet = wb[sheet_name]
                    table = list(sheet.iter_rows(values_only=True))
                    if not table or len(table) < 2: continue
                    
                    sheet_grade = extract_grade_from_text(sheet_name) or file_grade
                    prog_key, courses = process_table(table, dept_name, sheet_grade, page_num=s_idx)
                    if prog_key and courses:
                        if prog_key not in all_programs:
                            all_programs[prog_key] = {
                                "title": prog_key,
                                "department": dept_name,
                                "grade": courses[0]["grade"],
                                "file": fname,
                                "courses": []
                            }
                        all_programs[prog_key]["courses"].extend(courses)
                        total_slots += len(courses)
            except Exception as e:
                print(f"  [-] Excel Hatası ({fname}): {e}")

    js_content = "const ALL_PROGRAMS_DATA = " + json.dumps(all_programs, ensure_ascii=False, indent=2) + ";"
    with open("courses_data.js", "w", encoding="utf-8") as f:
        f.write(js_content)

    print(f"\n[✓] İŞLEM BAŞARIYLA TAMAMLANDI!")
    print(f"Toplam Program: {len(all_programs)}")
    print(f"Toplam Ayrıştırılan Ders Saati: {total_slots}")
    print("Oluşturulan dosya: courses_data.js")

if __name__ == "__main__":
    main()