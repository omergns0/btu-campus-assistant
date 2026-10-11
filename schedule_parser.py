# -*- coding: utf-8 -*-
"""
BTÜ ders programı ayrıştırıcısı (PDF / XLSX / XLS / DOCX).

Temel prensipler
----------------
1. Koordinat tabanlı okuma (PDF): pdfplumber tabloları sık sık parçalara
   bölüyor. Bu yüzden "tablo" yerine sayfadaki TÜM hücre kutularını alıp her
   hücreyi gün başlığının X aralığına ve saat etiketinin Y aralığına göre
   yerleştiriyoruz. Parçalanmış tablolar, birleşik (çok saatlik) hücreler ve
   aynı sayfadaki birden fazla sınıf tablosu böylece doğru okunur.
2. Sınıf (1-4) bilgisi ASLA ders kodundan çıkarılmaz. Kaynak sırası:
   tablo başlığı ("3. SINIF") > dosya meta verisi (site/ dosya adı) > tablo sırası.
   Ders kodundan gelen seviye sadece "level" alanına yazılır; seviye sınıftan
   farklıysa ders "elective" (seçmeli / alttan) olarak işaretlenir.
3. Saat formatları: 08:00, 8.00, 08.00 - 08.45, 8:15-9:00 ...
"""
import os
import re
import unicodedata

DAY_CANON = {
    "pazartesi": "Pazartesi", "sali": "Salı", "carsamba": "Çarşamba",
    "persembe": "Perşembe", "cuma": "Cuma", "cumartesi": "Cumartesi", "pazar": "Pazar",
    "monday": "Pazartesi", "tuesday": "Salı", "wednesday": "Çarşamba",
    "thursday": "Perşembe", "friday": "Cuma", "saturday": "Cumartesi",
}

TIME_RE = re.compile(r"(?<!\d)([01]?\d|2[0-3])\s*[:.]\s*([0-5]\d)(?!\d)")
CODE_RE = re.compile(r"(?<![A-Za-zÇĞİÖŞÜçğıöşü0-9])([A-ZÇĞİÖŞÜ]{2,5})\s?-?\s?(\d{3,4})(?![0-9])")
TITLE_RE = re.compile(r"\b(Prof\.?|Doç\.?|Dr\.?|Öğr\.?|Arş\.?|Ögr\.?|Okt\.?|Uzm\.?|Lect\.?|Assoc\.?|Asst\.?)\s", re.IGNORECASE)
ROOM_HINT_RE = re.compile(r"(\b(MS|YB)\s*/|\bblok\b|\bamfi|\blab|\bderslik|\bonline\b|\buzem\b|\bofis\b|kamp[üu]s|[A-Z]\s*[-/]\s*\d{2,3}\b|\bZ\d{2}\b|yer\.)", re.IGNORECASE)


# ---------------------------------------------------------------------------
# Yardımcılar
# ---------------------------------------------------------------------------
def norm(s):
    """Türkçe duyarlı, aksansız, küçük harfli normalizasyon."""
    if s is None:
        return ""
    s = str(s).replace("İ", "i").replace("I", "ı").lower()
    s = s.translate(str.maketrans("ıışçğüö", "iiscguo"))
    s = unicodedata.normalize("NFKD", s)
    return "".join(ch for ch in s if not unicodedata.combining(ch))


def clean(s):
    if s is None:
        return ""
    return re.sub(r"\s+", " ", str(s)).strip()


def match_day(text):
    """Hücre metni yalnızca bir gün adıysa kanonik gün adını döndürür."""
    if not text:
        return None
    t = re.sub(r"[^a-z]", "", norm(text))
    return DAY_CANON.get(t)


def first_hour(text):
    if not text:
        return None
    m = TIME_RE.search(str(text))
    if m:
        return int(m.group(1))
    m2 = re.search(r"(?<!\d)([01]?\d|2[0-3])\s*-\s*([01]?\d|2[0-3])(?!\d)", str(text))
    if m2:
        return int(m2.group(1))
    return None


GRADE_PATTERNS = [
    re.compile(r"(?<!\d)([1-6])\s*\.?\s*(?:sinif|class|grade|yariyil|yil)"),
    re.compile(r"(?:sinif|class|grade)\s*[:\-]?\s*([1-6])(?!\d)"),
    re.compile(r"(?<![a-z])(i{1,3}|iv)\s*\.\s*(?:sinif|yil)"),
]
ROMAN = {"i": 1, "ii": 2, "iii": 3, "iv": 4}


def find_grades(text):
    """Metindeki tüm sınıf ifadelerini (sırasıyla) döndürür."""
    t = norm(text)
    found = []
    for pat in GRADE_PATTERNS:
        for m in pat.finditer(t):
            g = m.group(1)
            g = ROMAN.get(g, g)
            try:
                g = int(g)
            except ValueError:
                continue
            if 1 <= g <= 6:
                found.append((m.start(), g))
    # Dikey/ters basılmış başlıklar: "FINIS .1" == "1. SINIF"
    for m in re.finditer(r"finis\s*\.?\s*([1-6])", t):
        found.append((m.start(), int(m.group(1))))
    found.sort()
    return [g for _, g in found]


def extract_grade(text):
    gs = find_grades(text)
    return gs[0] if gs else None


def course_level(code):
    """BLM0481 -> 4, BLM310 -> 3. Sadece bilgi amaçlı (sınıf ataması için KULLANILMAZ)."""
    m = re.search(r"(\d{3,4})$", code or "")
    if not m:
        return None
    digits = m.group(1)
    if len(digits) == 4:
        digits = digits[1:] if digits[0] == "0" else digits
    d = int(digits[0])
    return d if 1 <= d <= 6 else None


# ---------------------------------------------------------------------------
# Hücre içeriği -> ders kayıtları
# ---------------------------------------------------------------------------
def _split_courses(text):
    if "■" in text:
        return [p.strip() for p in text.split("■") if p.strip()]
    # Satır başında yeni bir ders kodu başlıyorsa böl
    lines = [l.strip() for l in text.split("\n") if l.strip()]
    chunks, cur = [], []
    for l in lines:
        if CODE_RE.match(l) and cur and CODE_RE.search(" ".join(cur)):
            chunks.append("\n".join(cur))
            cur = []
        cur.append(l)
    if cur:
        chunks.append("\n".join(cur))
    # Tek satırda birden fazla "KOD Ad T, ... Şube" kalıbı varsa yine böl
    out = []
    for ch in chunks:
        flat = clean(ch)
        starts = [m.start() for m in CODE_RE.finditer(flat)
                  if re.match(r"[^,]{2,90}?\s[TLU]\s*,", flat[m.end():] + " ,")]
        if len(starts) > 1 and len(re.findall(r"[ŞS]ube\s*\d", flat, re.I)) >= len(starts):
            for i, s in enumerate(starts):
                e = starts[i + 1] if i + 1 < len(starts) else len(flat)
                out.append(flat[s:e].strip(" ,"))
        else:
            out.append(ch)
    return out


def clean_classroom(room):
    if not room:
        return ""
    r = clean(room).strip(" ,-")
    # If it starts with a course code, it's not a classroom
    if re.search(r"^[A-ZÇĞİÖŞÜ]{2,5}\s*\d{3,4}", r):
        return ""
    if re.search(r"\b(mühendisli[gğ]i|program[ıi]|yarıyıl|ders\s*program)\b", r, re.I):
        return ""
    return r


def parse_course_cell(cell_text):
    """Bir hücredeki tüm dersleri ayrıştırır."""
    if not cell_text:
        return []
    raw_all = str(cell_text).replace("\r", "\n")
    if len(clean(raw_all)) < 4:
        return []
    results = []
    for chunk in _split_courses(raw_all):
        flat = clean(chunk)
        if len(flat) < 4:
            continue
        code_m = CODE_RE.search(flat)
        code = (code_m.group(1) + code_m.group(2)) if code_m else ""
        # Kod yoksa ve metin sadece öğretim üyesi/oda/not gibi görünüyorsa atla
        if not code:
            letters = re.sub(r"[^A-Za-zÇĞİÖŞÜçğıöşü]", "", flat)
            if len(letters) < 6 or TITLE_RE.match(flat) or re.fullmatch(r"(öğle|ogle)\s*(arası|arasi)?", flat, re.I):
                continue
        rec = {"raw": flat, "code": code, "name": "", "type": "", "section": "Şube 1",
               "instructor": "", "classroom": ""}

        sec_m = re.search(r"[ŞS]ube\s*:?\s*(\d+)", flat, re.IGNORECASE)
        grp_m = re.search(r"\b(?:Grup|Group)\s*:?\s*([A-Z0-9]+)", flat, re.IGNORECASE)
        if sec_m:
            rec["section"] = f"Şube {sec_m.group(1)}"
        elif grp_m:
            rec["section"] = f"Grup {grp_m.group(1)}"

        after = flat[code_m.end():].strip(" -:") if code_m else flat

        if "," in after and sec_m:
            # BTÜ standart formatı: "KOD Ad T, Bölüm Şube 1, Hoca, Derslik"
            parts = [p.strip() for p in after.split(",")]
            name = parts[0]
            sec_idx = next((i for i, p in enumerate(parts) if re.search(r"[ŞS]ube\s*\d", p, re.I)), None)
            if sec_idx is not None:
                rest = parts[sec_idx + 1:]
                instr_parts, room_parts = [], []
                for p in rest:
                    if ROOM_HINT_RE.search(p) or any(c in p for c in ['/', 'Blok', 'blok', 'Amfi', 'amfi', 'Lab', 'lab', 'Online', 'UZEM', 'Ofis', 'Yer.']):
                        room_parts.append(p)
                    elif not instr_parts and (TITLE_RE.search(p + " ") or any(ch.isupper() for ch in p)):
                        instr_parts.append(p)
                    else:
                        room_parts.append(p)
                rec["instructor"] = clean(", ".join(instr_parts))
                rec["classroom"] = clean(", ".join(room_parts))
            # Kod ile Şube arasındaki kısım: "Ad T, Bölüm" ya da "Ad T Bölüm"
            if sec_idx == 0:
                name = re.split(r"\s[TLU]\s", name + " ", maxsplit=1)[0]
        else:
            # Serbest format (İTBF vb.): satır bazlı sezgisel
            lines = [clean(l) for l in chunk.split("\n") if clean(l)]
            if code_m and lines:
                first = clean(lines[0])
                idx = first.find(code_m.group(0).strip())
                lines[0] = first[idx + len(code_m.group(0)):].strip(" -:") if idx >= 0 else first
            name_lines, instr, room = [], [], []
            for l in lines:
                if not l:
                    continue
                paren = re.findall(r"\(([^()]*)\)", l)
                core = re.sub(r"\([^()]*\)", " ", l).strip()
                for p in paren:
                    (room if ROOM_HINT_RE.search(p) else instr if TITLE_RE.search(p + " ") else name_lines if not instr else instr).append(p.strip())
                if not core:
                    continue
                if TITLE_RE.search(core + " "):
                    pos = TITLE_RE.search(core + " ").start()
                    if pos > 0 and not instr:
                        name_lines.append(core[:pos].strip())
                    instr.append(core[pos:].strip())
                elif ROOM_HINT_RE.search(core) and (instr or name_lines):
                    room.append(core)
                elif not instr:
                    name_lines.append(core)
                else:
                    room.append(core)
            name = clean(" ".join(name_lines))
            rec["instructor"] = clean(" ".join(instr))
            rec["classroom"] = clean(" ".join(room))

        name = clean(name).strip(" ,-")
        t_m = re.search(r"\s([TLU])$", name)
        if t_m:
            rec["type"] = {"T": "Teori", "L": "Lab", "U": "Uygulama"}[t_m.group(1)]
            name = name[:t_m.start()].strip()
        rec["name"] = (f"{name} ({rec['type']})" if rec["type"] and name else name) or code or flat[:60]
        rec["classroom"] = clean_classroom(rec["classroom"])
        rec["instructor"] = clean(rec["instructor"]).strip(" ,")
        rec["level"] = course_level(code)
        rec["elective"] = bool(re.search(r"se[cç]meli|elective", norm(flat)))
        results.append(rec)
    return results


# ---------------------------------------------------------------------------
# Ortak blok -> program dönüştürücü
# ---------------------------------------------------------------------------
class Block:
    """Bir sınıfa ait tek bir haftalık tablo."""

    def __init__(self, grade=None, explicit=False, origin=""):
        self.grade = grade
        self.explicit = explicit
        self.origin = origin
        self.entries = []          # (day, hour, cell_text)

    def add(self, day, hour, text):
        if text and clean(text):
            self.entries.append((day, hour, str(text)))


def _collect_courses(block):
    courses, seen = [], set()
    for day, hour, text in block.entries:
        for c in parse_course_cell(text):
            key = (day, hour, c["code"], c["section"], c["name"])
            if key in seen:
                continue
            seen.add(key)
            c = dict(c, day=day, hour=hour, time=f"{hour:02d}:00 - {hour:02d}:45")
            courses.append(c)
    return courses


# ---------------------------------------------------------------------------
# Matris (Excel / Word) ayrıştırıcı
# ---------------------------------------------------------------------------
def parse_matrix(rows, origin=""):
    """rows: list[list[str]] -> list[Block]"""
    rows = [[clean(c) if c is not None else "" for c in r] for r in rows]
    raw_rows = rows
    blocks = []
    cur = None
    cols = {}
    last_header_end = 0
    for r_idx, row in enumerate(raw_rows):
        found = {}
        for c_idx, cell in enumerate(row):
            d = match_day(cell)
            if d and d not in found:
                found[d] = c_idx
        if len(found) >= 3:
            # Yeni sınıf tablosu başlığı. Sınıf bilgisini başlık satırı ve üstündeki satırlarda ara.
            ctx = " ".join(" ".join(r) for r in raw_rows[last_header_end:r_idx + 1][-4:])
            g = extract_grade(ctx)
            cur = Block(g, explicit=g is not None, origin=f"{origin}#r{r_idx + 1}")
            blocks.append(cur)
            cols = found
            last_header_end = r_idx + 1
            continue
        if cur is None:
            continue
        min_col = min(cols.values())
        hour = None
        for cell in row[:max(min_col, 1)]:
            hour = first_hour(cell)
            if hour is not None:
                break
        if hour is None:
            # Bir sonraki tablonun başlığı satırda yazıyorsa (ör. "2. SINIF")
            continue
        for day, c_idx in cols.items():
            if c_idx < len(row):
                cur.add(day, hour, row[c_idx])
    return blocks


def _expand_merged_openpyxl(ws):
    grid = [[c for c in r] for r in ws.iter_rows(values_only=True)]
    for rng in ws.merged_cells.ranges:
        v = ws.cell(rng.min_row, rng.min_col).value
        for r in range(rng.min_row, rng.max_row + 1):
            for c in range(rng.min_col, rng.max_col + 1):
                if r - 1 < len(grid) and c - 1 < len(grid[r - 1]):
                    grid[r - 1][c - 1] = v
    return grid


def parse_excel(path):
    blocks = []
    ext = os.path.splitext(path)[1].lower()
    if ext == ".xls":
        import xlrd  # eski Excel
        wb = xlrd.open_workbook(path, formatting_info=True)
        for sh in wb.sheets():
            grid = [[sh.cell_value(r, c) for c in range(sh.ncols)] for r in range(sh.nrows)]
            for (r0, r1, c0, c1) in sh.merged_cells:
                v = grid[r0][c0]
                for r in range(r0, r1):
                    for c in range(c0, c1):
                        grid[r][c] = v
            bs = parse_matrix(grid, origin=f"{os.path.basename(path)}:{sh.name}")
            _apply_sheet_grade(bs, sh.name)
            blocks += bs
        return blocks
    import openpyxl
    wb = openpyxl.load_workbook(path, data_only=True)
    for ws in wb.worksheets:
        grid = _expand_merged_openpyxl(ws)
        bs = parse_matrix(grid, origin=f"{os.path.basename(path)}:{ws.title}")
        _apply_sheet_grade(bs, ws.title)
        blocks += bs
    return blocks


def _apply_sheet_grade(blocks, sheet_name):
    g = extract_grade(sheet_name)
    if g and len(blocks) == 1 and not blocks[0].explicit:
        blocks[0].grade, blocks[0].explicit = g, True


def parse_docx(path):
    import docx
    doc = docx.Document(path)
    blocks = []
    # Tablodan önceki paragraflar sınıf başlığı içerebilir
    body = doc.element.body
    last_text = ""
    t_idx = 0
    for child in body.iterchildren():
        tag = child.tag.split("}")[-1]
        if tag == "p":
            txt = "".join(n.text or "" for n in child.iter() if n.tag.endswith("}t"))
            if txt.strip():
                last_text = (last_text + " " + txt)[-300:]
        elif tag == "tbl":
            table = doc.tables[t_idx]
            t_idx += 1
            grid = [[cell.text for cell in row.cells] for row in table.rows]
            bs = parse_matrix(grid, origin=f"{os.path.basename(path)}#t{t_idx}")
            g = extract_grade(last_text)
            for b in bs:
                if not b.explicit and g:
                    b.grade, b.explicit = g, True
            blocks += bs
            last_text = ""
    return blocks


# ---------------------------------------------------------------------------
# PDF ayrıştırıcı (koordinat tabanlı)
# ---------------------------------------------------------------------------
def _chars_text(chars):
    from pdfplumber.utils import extract_text
    if not chars:
        return ""
    return extract_text(chars, x_tolerance=1.5, y_tolerance=3)


def _in_bbox(ch, bbox):
    cx = (ch["x0"] + ch["x1"]) / 2
    cy = (ch["top"] + ch["bottom"]) / 2
    return bbox[0] <= cx <= bbox[2] and bbox[1] <= cy <= bbox[3]


def _overlap(a0, a1, b0, b1):
    return max(0.0, min(a1, b1) - max(a0, b0))


def _region_grade(chars, x0, top, x1, bottom):
    """Bölgedeki karakterlerden sınıf bilgisi (dikey/ters yazılar dahil)."""
    sel = [c for c in chars if x0 <= (c["x0"] + c["x1"]) / 2 <= x1 and top <= (c["top"] + c["bottom"]) / 2 <= bottom]
    if not sel:
        return None
    upright = [c for c in sel if c.get("upright", True)]
    g = extract_grade(_chars_text(upright))
    if g:
        return g
    vertical = sorted([c for c in sel if not c.get("upright", True)], key=lambda c: (round(c["x0"]), c["top"]))
    s = "".join(c["text"] for c in vertical)
    return extract_grade(s) or extract_grade(s[::-1])


def _page_cells(page):
    cells = []
    try:
        tables = page.find_tables()
    except Exception:
        tables = []
    for t in tables:
        for bbox in t.cells:
            if bbox:
                cells.append(tuple(bbox))
    # yinelenenleri at
    uniq = []
    for b in sorted(set(cells), key=lambda b: (b[1], b[0])):
        if not any(abs(b[0] - u[0]) < 1 and abs(b[1] - u[1]) < 1 and abs(b[2] - u[2]) < 1 and abs(b[3] - u[3]) < 1 for u in uniq):
            uniq.append(b)
    return uniq


def _word_grid_cells(page, words):
    """Çizgisiz PDF'ler için yedek: kelimelerden sanal hücre üret."""
    cells = []
    for w in words:
        cells.append((w["x0"], w["top"], w["x1"], w["bottom"]))
    return cells


def _header_rows(cells_txt):
    """Gün başlığı hücrelerini satırlara grupla."""
    heads = [(b, match_day(t)) for b, t in cells_txt if match_day(t)]
    
    # Also check if any single cell contains multiple days (merged header cell)
    for b, t in cells_txt:
        if match_day(t): continue
        norm_t = norm(t)
        found_days = []
        for d_key, canon in DAY_CANON.items():
            if d_key in norm_t and canon not in found_days:
                found_days.append(canon)
        if len(found_days) >= 3:
            # Sort days by order of occurrence in norm_t
            found_days.sort(key=lambda d: norm_t.find(norm(d)))
            col_width = (b[2] - b[0]) / len(found_days)
            for i, d in enumerate(found_days):
                sub_b = (b[0] + i * col_width, b[1], b[0] + (i + 1) * col_width, b[3])
                heads.append((sub_b, d))

    rows = []
    for b, d in sorted(heads, key=lambda x: x[0][1]):
        for row in rows:
            if abs(row["top"] - b[1]) < 4:
                if d not in row["days"]:
                    row["days"][d] = b
                break
        else:
            rows.append({"top": b[1], "days": {d: b}})
    rows = [r for r in rows if len(r["days"]) >= 3]
    for r in rows:
        r["bottom"] = max(b[3] for b in r["days"].values())
        r["x0"] = min(b[0] for b in r["days"].values())
    return rows


def _time_rows(cells_txt, x_limit, top, bottom, words):
    rows = []
    for b, t in cells_txt:
        if b[2] <= x_limit + 3 and b[1] >= top - 2 and b[3] <= bottom + 2:
            h = first_hour(t)
            if h is not None:
                rows.append((b[1], b[3], h))
    if rows:
        # aynı saate ait iç içe hücreleri tekilleştir (en geniş olanı tut)
        best = {}
        for r in rows:
            if r[2] not in best or (r[1] - r[0]) > (best[r[2]][1] - best[r[2]][0]):
                best[r[2]] = r
        return sorted(best.values())
    # Yedek: sol sütundaki kelimelerden saat etiketleri
    left = [w for w in words if w["x1"] <= x_limit + 3 and top - 2 <= w["top"] <= bottom]
    lines = {}
    for w in left:
        key = round(w["top"] / 3)
        lines.setdefault(key, []).append(w)
    marks = []
    for key in sorted(lines):
        ws = sorted(lines[key], key=lambda w: w["x0"])
        txt = " ".join(w["text"] for w in ws)
        ts = TIME_RE.findall(txt)
        if not ts:
            continue
        h, mnt = int(ts[0][0]), int(ts[0][1])
        is_end = len(ts) == 1 and mnt >= 40
        marks.append((ws[0]["top"], h, is_end))
    starts = [(y, h) for y, h, e in marks if not e]
    out = []
    for i, (y, h) in enumerate(starts):
        y2 = starts[i + 1][0] - 1 if i + 1 < len(starts) else bottom
        out.append((y - 4, y2, h))
    return out


REVERSE_DAYS = {
    'pazartesi': 'Pazartesi', 'isetrazap': 'Pazartesi', 'isetraz': 'Pazartesi',
    'sali': 'Salı', 'ilas': 'Salı',
    'carsamba': 'Çarşamba', 'abmasrac': 'Çarşamba', 'abmasr': 'Çarşamba',
    'persembe': 'Perşembe', 'ebmesrep': 'Perşembe', 'ebmsrep': 'Perşembe',
    'cuma': 'Cuma', 'amuc': 'Cuma',
}

def match_any_day(text):
    if not text: return None
    clean_t = re.sub(r'[^a-zA-ZçğıöşüÇĞİÖŞÜ]', '', str(text)).lower()
    norm_t = norm(clean_t)
    norm_rev = norm(clean_t[::-1])
    for k, d in REVERSE_DAYS.items():
        if k in norm_t or k in norm_rev:
            return d
    return None

def parse_transposed_pdf(path):
    import pdfplumber
    blocks_by_grade = {}
    current_day = None
    prev_grade_cols = {}
    with pdfplumber.open(path) as pdf:
        for p_idx, page in enumerate(pdf.pages, 1):
            tables = page.find_tables()
            if not tables: continue
            grid = tables[0].extract()
            if not grid or len(grid) < 2: continue

            grade_cols = {}
            for row in grid[:2]:
                for c_idx, cell in enumerate(row):
                    g = extract_grade(cell)
                    if not g and cell:
                        m_yil = re.search(r'([1-6])\s*\.?\s*yil', norm(cell))
                        if m_yil: g = int(m_yil.group(1))
                    if g:
                        room_col = None
                        if c_idx + 1 < len(row) and re.search(r'derslik|oda', str(row[c_idx + 1]), re.I):
                            room_col = c_idx + 1
                        grade_cols[c_idx] = (g, room_col)
                if grade_cols: break

            if not grade_cols and prev_grade_cols:
                grade_cols = prev_grade_cols
            elif grade_cols:
                prev_grade_cols = grade_cols

            if not grade_cols: continue

            for row in grid:
                day_cand = match_any_day(row[0]) or (match_any_day(row[1]) if len(row) > 1 else None)
                if day_cand:
                    current_day = day_cand
                if not current_day: continue

                hour = None
                for cell in row[:4]:
                    h = first_hour(cell)
                    if h is not None:
                        hour = h
                        break
                if hour is None: continue

                for c_idx, (grade, room_col) in grade_cols.items():
                    if c_idx < len(row) and row[c_idx]:
                        val = row[c_idx].strip()
                        if not val or len(val) < 4: continue
                        room_val = row[room_col].strip() if (room_col and room_col < len(row) and row[room_col]) else ""
                        if grade not in blocks_by_grade:
                            blocks_by_grade[grade] = Block(grade, explicit=True, origin=f"{os.path.basename(path)}#g{grade}")
                        cell_txt = val
                        if room_val and not ROOM_HINT_RE.search(cell_txt):
                            cell_txt = f"{val}, {room_val}"
                        blocks_by_grade[grade].add(current_day, hour, cell_txt)

    return list(blocks_by_grade.values())


def parse_pdf(path):
    import pdfplumber
    blocks = []
    with pdfplumber.open(path) as pdf:
        prev_ctx = None   # önceki sayfadan devam eden tablo (sütunlar, blok)
        for p_idx, page in enumerate(pdf.pages, start=1):
            chars = page.chars
            upright = [c for c in chars if c.get("upright", True)]
            words = page.extract_words(keep_blank_chars=False)
            cells = _page_cells(page)
            used_word_fallback = False
            if not cells:
                cells = _word_grid_cells(page, words)
                used_word_fallback = True

            cells_txt = [(b, _chars_text([c for c in upright if _in_bbox(c, b)])) for b in cells]
            headers = _header_rows(cells_txt)

            segments = []   # (top, bottom, day_cols, block)
            first_top = headers[0]["top"] if headers else page.height
            if prev_ctx and first_top > 40:
                segments.append((0, first_top - 1, prev_ctx[0], prev_ctx[1], None))
            prev_bottom = 0
            for i, h in enumerate(headers):
                seg_bottom = headers[i + 1]["top"] - 1 if i + 1 < len(headers) else page.height
                day_cols = {d: (b[0], b[2]) for d, b in h["days"].items()}
                x_first = min(v[0] for v in day_cols.values())
                g = _region_grade(chars, 0, prev_bottom, page.width, h["bottom"] + 1)
                if g is None:
                    g = _region_grade(chars, 0, h["bottom"], x_first - 1, seg_bottom)
                blk = Block(g, explicit=g is not None, origin=f"{os.path.basename(path)}#p{p_idx}")
                blocks.append(blk)
                segments.append((h["bottom"], seg_bottom, day_cols, blk, h))
                prev_bottom = seg_bottom
            if segments:
                prev_ctx = (segments[-1][2], segments[-1][3])

            for top, bottom, day_cols, blk, h in segments:
                x_first = min(v[0] for v in day_cols.values())
                trows = _time_rows(cells_txt, x_first, top, bottom, words)
                if not trows:
                    continue
                if used_word_fallback:
                    # Kelime kelime yerleştir, sonra (gün, saat) bazında birleştir
                    buckets = {}
                    for w in words:
                        cy = (w["top"] + w["bottom"]) / 2
                        cx = (w["x0"] + w["x1"]) / 2
                        if not (top < cy <= bottom) or cx < x_first:
                            continue
                        day = min(day_cols, key=lambda d: abs((day_cols[d][0] + day_cols[d][1]) / 2 - cx))
                        hr = next((r[2] for r in trows if r[0] <= cy <= r[1]), None)
                        if hr is not None:
                            buckets.setdefault((day, hr), []).append(w)
                    for (day, hr), ws in buckets.items():
                        blk.add(day, hr, " ".join(x["text"] for x in sorted(ws, key=lambda w: (round(w["top"]), w["x0"]))))
                    continue
                for b, txt in cells_txt:
                    if not txt or match_day(txt):
                        continue
                    cy = (b[1] + b[3]) / 2
                    if not (top < cy <= bottom):
                        continue
                    if b[2] <= x_first + 3:
                        continue  # saat sütunu / sol başlık
                    days = [d for d, (x0, x1) in day_cols.items()
                            if _overlap(b[0], b[2], x0, x1) >= 0.5 * min(x1 - x0, b[2] - b[0])]
                    if not days:
                        continue
                    hours = [r[2] for r in trows
                             if _overlap(b[1], b[3], r[0], r[1]) >= 0.5 * min(r[1] - r[0], b[3] - b[1])]
                    for d in days:
                        for hr in hours:
                            blk.add(d, hr, txt)

    if not any(b.entries for b in blocks):
        transposed = parse_transposed_pdf(path)
        if any(b.entries for b in transposed):
            return transposed

    return blocks


# ---------------------------------------------------------------------------
# Dosya -> programlar
# ---------------------------------------------------------------------------
def parse_file(path):
    ext = os.path.splitext(path)[1].lower()
    if ext == ".pdf":
        return parse_pdf(path)
    if ext in (".xlsx", ".xlsm", ".xls"):
        return parse_excel(path)
    if ext == ".docx":
        return parse_docx(path)
    return []


def blocks_to_programs(blocks, dept_name, meta_grade=None):
    """
    Blokları sınıflara atar. Döner: ({grade: [course,...]}, uyarılar)
    Öncelik: açık başlık > meta (dosya/site) sınıfı > blok sırası.
    """
    warnings = []
    blocks = [b for b in blocks if b.entries]
    if not blocks:
        return {}, ["Dosyada okunabilir ders tablosu bulunamadı"]
    explicit = [b for b in blocks if b.explicit]
    for i, b in enumerate(blocks):
        if b.explicit:
            if meta_grade and len(blocks) == 1 and b.grade != meta_grade:
                warnings.append(f"Başlıkta {b.grade}. sınıf yazıyor, dosya/site {meta_grade}. sınıf diyor -> başlık esas alındı")
            continue
        if meta_grade and not explicit:
            b.grade = meta_grade
        else:
            b.grade = i + 1
            warnings.append(f"{b.origin}: sınıf başlığı bulunamadı, tablo sırasından {b.grade}. sınıf varsayıldı")
    out = {}
    for b in blocks:
        cs = _collect_courses(b)
        for c in cs:
            c["department"] = dept_name
            c["grade"] = b.grade
            lv = c.get("level")
            if lv and lv != b.grade:
                c["elective"] = True   # üst/alt sınıf kodu: seçmeli ya da alttan
        out.setdefault(b.grade, []).extend(cs)
    return out, warnings
