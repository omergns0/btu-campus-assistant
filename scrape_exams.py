# -*- coding: utf-8 -*-
"""
BTU Sınav Takvimi Tarayıcısı (BTU Exam Schedule Scraper) - V3 (2026-2027 Güz)
- 2026-2027 Güz Yarıyılı Ara Sınav (Vize) Programlarını çeker.
- Kesinleşmiş (FINAL), Taslak (DRAFT) ve Henüz Yayınlanmayan (PENDING) durumlarını ayrıştırır.
- Tüm fakülteler (MDBF, ITBF, MTF, OF, DF) ve 32 bölüm taranır.
- PDF ve Excel (.xlsx) sınav tabloları otomatik ayrıştırılır.
- JSON ve JS dosyaları üretilir.
"""
import os
import sys
import re
import io
import json
import shutil
import requests
from bs4 import BeautifulSoup
from urllib.parse import urljoin, unquote
from datetime import datetime

# Windows konsolunda UTF-8 desteği sağla
try:
    sys.stdout.reconfigure(encoding='utf-8')
except Exception:
    pass

try:
    import pdfplumber
except ImportError:
    pdfplumber = None

try:
    import openpyxl
except ImportError:
    openpyxl = None

from departments import DEPARTMENTS

HEADERS = {
    "User-Agent": "Mozilla/5.0 (Windows NT 10.0; Win64; x64) AppleWebKit/537.36 (KHTML, like Gecko) Chrome/120.0.0.0 Safari/537.36"
}

FACULTY_MAP = {
    "bilgisayar": "Mühendislik ve Doğa Bilimleri Fakültesi",
    "biyomuh": "Mühendislik ve Doğa Bilimleri Fakültesi",
    "cevre": "Mühendislik ve Doğa Bilimleri Fakültesi",
    "insaat": "Mühendislik ve Doğa Bilimleri Fakültesi",
    "gida": "Mühendislik ve Doğa Bilimleri Fakültesi",
    "fizik": "Mühendislik ve Doğa Bilimleri Fakültesi",
    "endustri": "Mühendislik ve Doğa Bilimleri Fakültesi",
    "elektrik": "Mühendislik ve Doğa Bilimleri Fakültesi",
    "kimya": "Mühendislik ve Doğa Bilimleri Fakültesi",
    "kimyamuh": "Mühendislik ve Doğa Bilimleri Fakültesi",
    "makine": "Mühendislik ve Doğa Bilimleri Fakültesi",
    "matematik": "Mühendislik ve Doğa Bilimleri Fakültesi",
    "mekatronik": "Mühendislik ve Doğa Bilimleri Fakültesi",
    "mekatroniking": "Mühendislik ve Doğa Bilimleri Fakültesi",
    "metalurji": "Mühendislik ve Doğa Bilimleri Fakültesi",
    "polimer": "Mühendislik ve Doğa Bilimleri Fakültesi",
    "veribilimi": "Mühendislik ve Doğa Bilimleri Fakültesi",
    "yapayzeka": "Mühendislik ve Doğa Bilimleri Fakültesi",
    "sosyoloji": "İnsan ve Toplum Bilimleri Fakültesi",
    "ui": "İnsan ve Toplum Bilimleri Fakültesi",
    "ui_ing": "İnsan ve Toplum Bilimleri Fakültesi",
    "utl": "İnsan ve Toplum Bilimleri Fakültesi",
    "isletme": "İnsan ve Toplum Bilimleri Fakültesi",
    "imtb": "İnsan ve Toplum Bilimleri Fakültesi",
    "psikoloji": "İnsan ve Toplum Bilimleri Fakültesi",
    "mimarlik": "Mimarlık ve Tasarım Fakültesi",
    "sehir": "Mimarlık ve Tasarım Fakültesi",
    "peyzaj": "Mimarlık ve Tasarım Fakültesi",
    "ormanendustri": "Orman Fakültesi",
    "orman": "Orman Fakültesi",
    "diy": "Denizcilik Fakültesi",
    "gigm": "Denizcilik Fakültesi"
}

EXAM_TYPE_KEYWORDS = ["ara sınav", "arasınav", "vize", "midterm"]
YEAR_KEYWORDS = ["2026-2027", "2026/2027", "2026 - 2027", "26-27", "26 - 27", "2026 güz", "26 güz"]

STRICT_EXCLUDE_WORDS = [
    "2025-2026", "2025 - 2026", "25-26", "25 - 26", "2024-2025", "2023-2024",
    "bütünleme", "butunleme", "final", "yarıyıl sonu", "yariyil sonu",
    "tek ders", "ek sınav", "yeterlik", "muafiyet", "yatay geçiş",
    "çift anadal", "doktora bilimsel", "mülakat", "personel", "öğrenci işleri staj",
    "bahar"
]

NON_EXAM_FILE_KEYWORDS = [
    "staj", "yönerge", "yonerge", "rapor", "akış", "akis", "ücret", "ucret", 
    "fonu", "anket", "başvuru", "basvuru", "katkı payı", "katki payi"
]

def clean_text(text):
    if not text:
        return ""
    return re.sub(r"\s+", " ", str(text)).strip()

def extract_date(text):
    m = re.search(r"(\d{1,2}[./-]\d{1,2}[./-](?:\d{4}|\d{2}))", text)
    if m:
        val = m.group(1).replace("/", ".").replace("-", ".")
        parts = val.split(".")
        if len(parts) == 3:
            day = int(parts[0])
            month = int(parts[1])
            year = parts[2]
            if len(year) == 2:
                year = "20" + year
            return f"{day:02d}.{month:02d}.{year}"
    return None

def is_valid_2026_midterm(text, url):
    """2026-2027 Güz Ara Sınav / Vize duyurularını doğrular."""
    combined = (clean_text(text) + " " + unquote(url)).lower()

    for ex in STRICT_EXCLUDE_WORDS:
        if ex in combined:
            return False

    # 'revize' kelimesi içindeki 'vize'yi elemek için temizle
    clean_combined = re.sub(r"\brevize\b", "", combined)
    has_exam_type = bool(re.search(r"\b(ara\s*s[ıi]nav|aras[ıi]nav|vize|midterm)\b", clean_combined))
    if not has_exam_type:
        return False

    # Sadece ders programı olanları (ara sınav / vize geçmeyen) reddet
    if "ders program" in combined and not bool(re.search(r"\b(ara\s*s[ıi]nav|aras[ıi]nav|vize|midterm)\b", clean_combined)):
        return False

    has_year = any(yw in combined for yw in YEAR_KEYWORDS)
    is_recent_2026 = "10.2026" in combined or "09.2026" in combined or "11.2026" in combined or "2026" in combined

    return has_year or is_recent_2026

def parse_pdf_exam_table(pdf_bytes):
    """Farklı fakültelerin sınav PDF'lerindeki ders tablosunu ayrıştırır."""
    if not pdfplumber:
        return []
    
    exam_items = []
    try:
        with pdfplumber.open(io.BytesIO(pdf_bytes)) as pdf:
            current_grade = ""
            for page in pdf.pages:
                tables = page.extract_tables()
                for table in tables:
                    if not table or len(table) < 2:
                        continue
                    
                    for row in table:
                        clean_row = [clean_text(c) for c in row if c]
                        row_str = " ".join(clean_row)
                        
                        # Sınıf başlıklarını yakala (1. Sınıf, 2.SINIF vb.)
                        grade_match = re.search(r"\b([1-4])\s*[\.\-]?\s*s[ıi]n[ıi]f", row_str, re.I)
                        if grade_match:
                            current_grade = grade_match.group(1)
                            # Eğer satır sadece sınıf başlığıysa ders olarak işleme
                            if len(clean_row) <= 2:
                                continue

                        # Başlık satırlarını atla
                        if any(h in row_str.lower() for h in ["ders kodu", "d.kodu", "ders adı", "oturum", "derslik 1", "sınav tarihi"]):
                            continue
                        if "vize program" in row_str.lower() or "ara sınav program" in row_str.lower():
                            continue

                        # 1. Standart MDBF Oturum formatı: "03/11/2026 13.00-14.45"
                        session_val = ""
                        for cell in row:
                            c_str = clean_text(cell)
                            if re.search(r"\d{1,2}[./-]\d{1,2}[./-](?:\d{4}|\d{2})", c_str):
                                session_val = c_str
                                break
                        
                        if session_val:
                            # Tarih ve saat ayrıştırma
                            date_match = re.search(r"(\d{1,2})[./-](\d{1,2})[./-](\d{4}|\d{2})", session_val)
                            time_match = re.search(r"(\d{1,2}[:.]\d{2})\s*[-–]\s*(\d{1,2}[:.]\d{2})", session_val)
                            
                            exam_date = ""
                            if date_match:
                                y = date_match.group(3)
                                if len(y) == 2:
                                    y = "20" + y
                                exam_date = f"{int(date_match.group(1)):02d}.{int(date_match.group(2)):02d}.{y}"
                                
                            exam_time = ""
                            if time_match:
                                exam_time = f"{time_match.group(1).replace('.', ':')} - {time_match.group(2).replace('.', ':')}"
                            else:
                                single_t = re.search(r"\b(\d{1,2}[:.]\d{2})\b", session_val)
                                if single_t:
                                    exam_time = single_t.group(1).replace(".", ":")

                            # Tablo formatını tespit et
                            # MDBF format: [Index, Ders Adı, D.Kodu, Şb, Bölüm, Sınıf, Hoca, ...]
                            course_code = ""
                            course_name = ""
                            instructor = ""
                            grade = current_grade
                            classrooms = []

                            # Kod tespiti (örn: MAT0101, BLM0111, PSİ0101, PM101)
                            for c in clean_row:
                                if re.search(r"^[A-ZÇĞİÖŞÜ]{2,4}\s*\d{3,4}$", c):
                                    course_code = c
                                    break

                            if len(row) >= 7 and "MD" in str(row[0] or ""):
                                # Klasik MDBF tablosu
                                course_name = clean_text(row[1])
                                course_code = clean_text(row[2])
                                grade = clean_text(row[5]) or grade
                                instructor = clean_text(row[6])
                                
                                session_idx = -1
                                for idx, cell in enumerate(row):
                                    if cell == session_val:
                                        session_idx = idx
                                        break
                                if session_idx != -1:
                                    for c in row[session_idx+1:]:
                                        cl = clean_text(c)
                                        if cl and len(cl) >= 2 and cl.upper() not in ["M", "MS", "YB", "VAR", "YOK", "DERSLİK"]:
                                            classrooms.append(cl)
                            else:
                                # ITBF veya MTF formatı
                                # [Hoca, Kod, Ders, Tarih, Saat, Salon] veya [Kod, Ders, Tarih, Saat, Salon]
                                for idx, c in enumerate(row):
                                    c_clean = clean_text(c)
                                    if c_clean == course_code:
                                        if idx > 0 and len(clean_text(row[idx-1])) > 5:
                                            instructor = clean_text(row[idx-1])
                                        if idx + 1 < len(row):
                                            course_name = clean_text(row[idx+1])
                                        break
                                if not course_name:
                                    for c in clean_row:
                                        if c != course_code and c != session_val and len(c) > 4 and not re.search(r"\d{2}[./-]", c):
                                            course_name = c
                                            break
                                classrooms_str = clean_row[-1] if len(clean_row) > 3 and clean_row[-1] != session_val else ""
                                if classrooms_str:
                                    classrooms = [classrooms_str]

                            if course_name or course_code:
                                exam_items.append({
                                    "code": course_code,
                                    "name": course_name,
                                    "date": exam_date,
                                    "time": exam_time,
                                    "instructor": instructor,
                                    "grade": grade,
                                    "classrooms": ", ".join(classrooms) if classrooms else "Derslikler Açıklanacak",
                                    "raw_session": session_val
                                })
    except Exception as e:
        print(f"[-] PDF ayrıştırma hatası: {e}")
        
    return exam_items

def parse_xlsx_exam_table(xlsx_bytes):
    """Excel (.xlsx) sınav tablosunu ayrıştırır."""
    if not openpyxl:
        return []
    
    exam_items = []
    try:
        wb = openpyxl.load_workbook(io.BytesIO(xlsx_bytes), data_only=True)
        sheet = wb.active

        # Sütunlar başlık satırındaki adlara göre bulunur; dosyadan dosyaya sütun sırası
        # değişebiliyor (ör. 2026 kesinleşmiş MDBF dosyasında "Tür" sütunu yok).
        header = [clean_text(sheet.cell(1, c).value).lower() for c in range(1, 25)]

        def hcol(*names, default=None):
            for i, h in enumerate(header):
                if h and any(n in h for n in names):
                    return i
            return default

        i_name = hcol("ders adı", "ders adi", default=1)
        i_code = hcol("kodu", default=2)
        i_dept = hcol("bölüm", "bolum")
        i_grade = hcol("sınıf", "sinif", default=5)
        i_inst = hcol("öğretim", "ogretim", default=6)
        # "Birleşeceği Dersler (Aynı Oturum)" sütunu da "oturum" içerdiği için tam eşleşme önce aranır
        i_session = header.index("oturum") if "oturum" in header else hcol("oturum", default=11)
        room_cols = [i for i, h in enumerate(header) if "derslik" in h] or list(range(12, 19))

        for row_idx in range(2, sheet.max_row + 1):
            row = [clean_text(sheet.cell(row_idx, col_idx).value) for col_idx in range(1, 25)]
            if not row or not any(row):
                continue

            course_name = row[i_name] if i_name < len(row) else ""
            course_code = row[i_code] if i_code < len(row) else ""
            grade = row[i_grade] if i_grade < len(row) else ""
            instructor = row[i_inst] if i_inst < len(row) else ""
            department = row[i_dept] if i_dept is not None and i_dept < len(row) else ""
            session_val = row[i_session] if i_session < len(row) else ""

            if not session_val:
                for c in row:
                    if re.search(r"\d{1,2}[./-]\d{1,2}[./-]\d{4}", c):
                        session_val = c
                        break
            
            if not session_val or (not course_name and not course_code):
                continue
                
            date_match = re.search(r"(\d{1,2})[./-](\d{1,2})[./-](\d{4})", session_val)
            time_match = re.search(r"(\d{1,2}[:.]\d{2})\s*[-–]\s*(\d{1,2}[:.]\d{2})", session_val)
            
            exam_date = ""
            if date_match:
                exam_date = f"{int(date_match.group(1)):02d}.{int(date_match.group(2)):02d}.{date_match.group(3)}"
                
            exam_time = ""
            if time_match:
                exam_time = f"{time_match.group(1).replace('.', ':')} - {time_match.group(2).replace('.', ':')}"
                
            classrooms = []
            for i in room_cols:
                c = row[i] if i < len(row) else ""
                if c and len(c) >= 2 and c.upper() not in ["M", "MS", "YB", "VAR", "YOK"]:
                    classrooms.append(c)

            item = {
                "code": course_code,
                "name": course_name,
                "date": exam_date,
                "time": exam_time,
                "instructor": instructor,
                "grade": grade,
                "classrooms": ", ".join(classrooms) if classrooms else "Derslikler Açıklanacak",
                "raw_session": session_val
            }
            if department:
                item["department"] = department
            exam_items.append(item)
    except Exception as e:
        print(f"[-] Excel ayrıştırma hatası: {e}")
        
    return exam_items

def parse_file_exams(file_url):
    """Dosya türüne göre (PDF veya Excel) sınav derslerini çeker."""
    lower = file_url.lower()
    try:
        r = requests.get(file_url, headers=HEADERS, timeout=12)
        if r.status_code != 200:
            return []
        if lower.endswith(".pdf"):
            return parse_pdf_exam_table(r.content)
        elif lower.endswith(".xlsx") or lower.endswith(".xls"):
            return parse_xlsx_exam_table(r.content)
    except Exception as e:
        print(f"[-] Dosya indirme/ayrıştırma hatası ({file_url}): {e}")
    return []

def is_exam_file(file_name, file_url):
    """Sınavla ilgisiz staj yönergesi, iş akışı gibi dosyaları eler."""
    comb = (clean_text(file_name) + " " + unquote(file_url)).lower()
    for kw in NON_EXAM_FILE_KEYWORDS:
        if kw in comb:
            return False
    return True

def parse_html_exam_table(soup):
    """Programı PDF yerine duyuru sayfasında tablo olarak veren bölümler (ör. Fizik)."""
    items = []
    for table in soup.find_all("table"):
        rows = table.find_all("tr")
        if len(rows) < 2:
            continue
        header = [clean_text(c.get_text(" ")).lower() for c in rows[0].find_all(["th", "td"])]

        def col(*names):
            for i, h in enumerate(header):
                if any(n in h for n in names):
                    return i
            return None

        i_date, i_time = col("tarih"), col("saat")
        i_name, i_code = col("ders adı", "ders adi"), col("ders kodu", "kod")
        i_grade, i_inst, i_room = col("sınıf", "sinif"), col("öğretim", "ogretim", "hoca"), col("derslik")
        if i_date is None or i_code is None:
            continue
        for r in rows[1:]:
            cells = [clean_text(c.get_text(" ")) for c in r.find_all(["td", "th"])]
            if len(cells) <= max(i for i in [i_date, i_code] if i is not None):
                continue
            get = lambda i: cells[i] if i is not None and i < len(cells) else ""
            date = extract_date(get(i_date))
            code = get(i_code).replace(" ", "")
            if not date or not code:
                continue
            t = re.findall(r"(\d{1,2})[.:](\d{2})", get(i_time))
            time = f"{int(t[0][0]):02d}:{t[0][1]} - {int(t[1][0]):02d}:{t[1][1]}" if len(t) >= 2 else get(i_time)
            grade = re.search(r"\d", get(i_grade))
            items.append({
                "code": code,
                "name": get(i_name),
                "date": date,
                "time": time,
                "instructor": get(i_inst),
                "grade": grade.group(0) if grade else "",
                "classrooms": get(i_room),
                "raw_session": f"{get(i_date)} {get(i_time)}"
            })
    return items


def find_pdf_links_in_page(url, parent_title="", _depth=0):
    try:
        r = requests.get(url, headers=HEADERS, timeout=10)
        if r.status_code != 200:
            return []
        soup = BeautifulSoup(r.text, "html.parser")
        files = []
        # Dosya yerine sayfa içi tablo (ör. Fizik bölümü)
        table_items = parse_html_exam_table(soup)
        if table_items:
            files.append({
                "title": "2026-2027 Güz Ara Sınav Programı (Duyuru sayfası)",
                "url": url,
                "ext": "html",
                "exam_items": table_items
            })
        # "Programa ulaşmak için tıklayınız" gibi başka bir duyuruya yönlendirmeler
        # (ör. Kimya Mühendisliği → MDBF ortak kesinleşmiş program) bir kez takip edilir
        if _depth == 0:
            for a in soup.find_all("a", href=True):
                text = clean_text(a.get_text()).lower()
                target = urljoin(url, a["href"].strip())
                if ("tıkla" in text or "tikla" in text) and "/duyuru/detay/" in target and target != url:
                    linked = find_pdf_links_in_page(target, parent_title, _depth=1)
                    if linked:
                        return files + linked
        for a in soup.find_all("a", href=True):
            href = a["href"].strip()
            full_url = urljoin(url, href)
            lower_url = full_url.lower()
            if any(lower_url.endswith(ext) for ext in [".pdf", ".xlsx", ".xls"]):
                raw_name = clean_text(a.get_text())
                if not is_exam_file(raw_name, full_url):
                    continue
                
                if not raw_name or len(raw_name) < 4 or any(w in raw_name.lower() for w in ["tıklayınız", "indir", "buradan", "link", "ek"]):
                    ext_label = "PDF" if lower_url.endswith(".pdf") else "Excel"
                    title = f"2026-2027 Güz Ara Sınav Programı ({ext_label})"
                else:
                    title = raw_name
                
                parsed_exams = parse_file_exams(full_url)
                
                files.append({
                    "title": title,
                    "url": full_url,
                    "ext": full_url.split(".")[-1].split("?")[0].lower(),
                    "exam_items": parsed_exams
                })
        return files
    except Exception:
        return []

def classify_status(announcements):
    """
    Duyuruyu 'FINAL' (Kesinleşmiş), 'DRAFT' (Taslak) veya 'PENDING' (Yayınlanmadı) olarak sınıflandırır.
    """
    if not announcements:
        return "PENDING", "Henüz Yayınlanmadı"
        
    latest_title = clean_text(announcements[0].get("title", "")).lower()
    
    # Kesinleşmiş belirteçleri
    is_explicit_final = any(w in latest_title for w in ["kesinleş", "kesinles", "nihai", "kesin program", "kesinleştirilmiş", "kesinlestirilmis"])
    # Taslak belirteçleri
    is_draft = any(w in latest_title for w in ["taslak", "taslağı", "taslagi", "draft"])
    
    if is_explicit_final:
        return "FINAL", "Kesinleşmiş Program"
    elif is_draft or "mdbf" in latest_title:
        return "DRAFT", "Taslak Program"
    else:
        # Başlığında taslak veya mdbf taslağı geçmeyen resmi ara sınav programları kesinleşmiş kabul edilir
        return "FINAL", "Kesinleşmiş Program"

# ITBF Bölümleri ve Özel PDF Eşleştirmeleri
ITBF_ANN_URL = "https://itbf.btu.edu.tr/tr/duyuru/detay/55081/2026-2027-güz-dönemi-kesinleştirilmiş-ara-sınav-programları"
ITBF_DEPT_FILES = {
    "isletme": {
        "title": "2026-2027 Güz Kesinleştirilmiş Ara Sınav Programı (PDF)",
        "url": "https://depo.btu.edu.tr/img/duyurular//1791550760_2ffde9f1520f248374d8.pdf"
    },
    "psikoloji": {
        "title": "2026-2027 Güz Kesinleştirilmiş Ara Sınav Programı (PDF)",
        "url": "https://depo.btu.edu.tr/img/duyurular//1791887068_53929105d9fddf6e9ea6.pdf"
    },
    "sosyoloji": {
        "title": "2026-2027 Güz Kesinleştirilmiş Ara Sınav Programı (PDF)",
        "url": "https://depo.btu.edu.tr/img/duyurular//1791550803_55c14f2d1aefc44189c6.pdf"
    },
    "ui_ing": {
        "title": "2026-2027 Güz Kesinleştirilmiş Ara Sınav Programı (%30 İngilizce) (PDF)",
        "url": "https://depo.btu.edu.tr/img/duyurular//1791872790_c03f7e287d625ea3005d.pdf"
    },
    "ui": {
        "title": "2026-2027 Güz Kesinleştirilmiş Ara Sınav Programı (%100 Türkçe) (PDF)",
        "url": "https://depo.btu.edu.tr/img/duyurular//1791872843_ea867c47b98f11d7f3c1.pdf"
    },
    "utl": {
        "title": "2026-2027 Güz Kesinleştirilmiş Ara Sınav Programı (PDF)",
        "url": "https://depo.btu.edu.tr/img/duyurular//1791552057_347890338910065aad03.pdf"
    },
    "imtb": {
        "title": "2026-2027 Güz Kesinleştirilmiş Ara Sınav Programı (PDF)",
        "url": "https://depo.btu.edu.tr/img/duyurular//1791550880_826cb7aea068cb67b76e.pdf"
    }
}

# Özel tanımlı güncel duyurular
SPECIFIC_DEPARTMENT_OVERRIDES = {
    "sehir": {
        "title": "2026-2027 Güz Dönemi Ara Sınav Programı",
        "date": "14.10.2026",
        "url": "https://mtf.btu.edu.tr/tr/sehir/duyuru/detay/55133/2026-2027-güz-dönemi-ara-sınav-programı",
        "file_title": "2026-2027 Güz Ara Sınav Programı (PDF)",
        "file_url": "https://depo.btu.edu.tr/img/duyurular//1791962145_8f657fd8987998dea1ee.pdf"
    },
    "peyzaj": {
        "title": "2026-2027 GÜZ DÖNEMİ ARA SINAV PROGRAMI",
        "date": "09.10.2026",
        "url": "https://mtf.btu.edu.tr/tr/peyzaj/duyuru/detay/55074/2026-2027-güz-dönemi̇-ara-sinav-programi",
        "file_title": "2026-2027 Güz Ara Sınav Programı (PDF)",
        "file_url": "https://depo.btu.edu.tr/img/duyurular//1791531746_b3741f60f104a0713c01.pdf"
    }
}
# Not: Metalurji, Polimer ve Mekatronik için eskiden burada sabit taslak duyurular vardı;
# bu yüzden kesinleşmiş programlar yayınlansa da görülmüyordu. Artık canlı taranıyorlar.

# Kendi duyuru sayfası olmayan, programını başka bir bölümle paylaşan bölümler
SHARED_DEPARTMENTS = {
    "mekatroniking": "mekatronik",
}

# Lisans öğrencileri için olmayan sınav duyuruları
GRADUATE_WORDS = ["lisans üstü", "lisansüstü", "lisansustu", "yüksek lisans", "yuksek lisans", "doktora", "tezsiz"]


def _fold(s):
    s = (s or "").lower()
    for a, b in (("ı", "i"), ("İ", "i"), ("i̇", "i"), ("ğ", "g"), ("ü", "u"), ("ş", "s"), ("ö", "o"), ("ç", "c")):
        s = s.replace(a, b)
    return re.sub(r"\s+", " ", s).strip()


def filter_for_department(items, dept_name):
    """Fakülte ortak dosyasındaki dersleri (ör. MDBF'nin 382 dersi) bölüme göre süzer.
    Dosyada bölüm bilgisi yoksa ya da hiçbir satır eşleşmezse liste olduğu gibi kalır."""
    target = _fold(re.sub(r"\(.*?\)", "", dept_name))
    if not target:
        return items
    # Tam eşleşme: "Kimya" bölümünün dersleri "Kimya Mühendisliği"ne karışmasın
    own = [it for it in items if it.get("department") and
           _fold(re.sub(r"\(.*?\)", "", it["department"])) == target]
    return own or items


# Fakültenin tüm bölümleri için ortak kesinleşmiş program; bir bölümün kendi dosyası
# okunamazsa (ör. Kimya'nın döndürülmüş PDF'i) dersleri buradan bölüme göre süzülür.
FACULTY_COMMON_FINAL = {
    "Mühendislik ve Doğa Bilimleri Fakültesi":
        "https://mdbf.btu.edu.tr/tr/duyuru/detay/55175/26-27-güz-mdbf-kesinleşmiş-arasınav-programı",
}
_common_cache = {}


def common_faculty_items(faculty, dept_name):
    url = FACULTY_COMMON_FINAL.get(faculty)
    if not url:
        return [], []
    if url not in _common_cache:
        _common_cache[url] = find_pdf_links_in_page(url)
    files = _common_cache[url]
    items = [it for f in files if f.get("ext") in ("xlsx", "xls") for it in f.get("exam_items", [])]
    own = filter_for_department(items, dept_name)
    if own is items:  # hiçbir satır bu bölüme ait değil
        return [], []
    return own, [f for f in files if f.get("ext") in ("xlsx", "xls")]


def announcement_sort_key(ann):
    """Sınav programı duyurularını öne, aynı grupta en yeni tarihliyi en başa alır."""
    title = (ann.get("title") or "").lower()
    is_program = "program" in title
    d = ann.get("date") or ""
    m = re.match(r"(\d{2})\.(\d{2})\.(\d{4})", d)
    date_key = (m.group(3), m.group(2), m.group(1)) if m else ("0000", "00", "00")
    return (is_program, date_key)

def scrape_department_exams(dept_key, dept_info):
    url = dept_info.get("url")
    dept_name = dept_info.get("name", dept_key)
    faculty_name = FACULTY_MAP.get(dept_key, "BTÜ Fakültesi")
    
    result = {
        "key": dept_key,
        "name": dept_name,
        "faculty": faculty_name,
        "url": url,
        "has_announcement": False,
        "status": "PENDING",
        "status_label": "Henüz Yayınlanmadı",
        "announcements": [],
        "last_checked": datetime.now().strftime("%d.%m.%Y %H:%M")
    }
    
    # 1. Özel tanımlı ITBF bölümleri
    if dept_key in ITBF_DEPT_FILES:
        f_info = ITBF_DEPT_FILES[dept_key]
        parsed_items = parse_file_exams(f_info["url"])
        
        file_obj = {
            "title": f_info["title"],
            "url": f_info["url"],
            "ext": "pdf",
            "exam_items": parsed_items
        }
        
        result["has_announcement"] = True
        result["status"] = "FINAL"
        result["status_label"] = "Kesinleşmiş Program"
        result["announcements"].append({
            "title": "2026-2027 Güz Dönemi Kesinleştirilmiş Ara Sınav Programları",
            "date": "09.10.2026",
            "url": ITBF_ANN_URL,
            "files": [file_obj],
            "exam_items": parsed_items
        })
        return result

    # 2. Özel tanımlı güncel duyuru geçersiz kılmaları
    if dept_key in SPECIFIC_DEPARTMENT_OVERRIDES:
        ov = SPECIFIC_DEPARTMENT_OVERRIDES[dept_key]
        parsed_items = parse_file_exams(ov["file_url"])
        
        file_obj = {
            "title": ov["file_title"],
            "url": ov["file_url"],
            "ext": ov["file_url"].split(".")[-1].split("?")[0].lower(),
            "exam_items": parsed_items
        }
        
        ann_obj = {
            "title": ov["title"],
            "date": ov["date"],
            "url": ov["url"],
            "files": [file_obj],
            "exam_items": parsed_items
        }
        
        result["has_announcement"] = True
        status, label = classify_status([ann_obj])
        result["status"] = status
        result["status_label"] = label
        result["announcements"].append(ann_obj)
        return result

    if not url:
        return result

    # 3. Canlı sayfa taraması
    try:
        r = requests.get(url, headers=HEADERS, timeout=12)
        if r.status_code != 200:
            return result
        
        soup = BeautifulSoup(r.text, "html.parser")
        candidate_links = []
        duyuru_list_urls = []
        
        for a in soup.find_all("a", href=True):
            href = a["href"].strip()
            text = clean_text(a.get_text())
            full_url = urljoin(url, href)
            
            if "/duyuru/birim/" in href or "/duyurular" in href:
                duyuru_list_urls.append(full_url)
            
            if is_valid_2026_midterm(text, full_url):
                candidate_links.append((text, full_url))
        
        if not candidate_links and duyuru_list_urls:
            archive_url = duyuru_list_urls[0]
            try:
                r_arch = requests.get(archive_url, headers=HEADERS, timeout=10)
                if r_arch.status_code == 200:
                    soup_arch = BeautifulSoup(r_arch.text, "html.parser")
                    for a in soup_arch.find_all("a", href=True):
                        href = a["href"].strip()
                        text = clean_text(a.get_text())
                        full_url = urljoin(archive_url, href)
                        if is_valid_2026_midterm(text, full_url):
                            candidate_links.append((text, full_url))
            except Exception:
                pass

        seen_urls = set()
        for title, item_url in candidate_links:
            if item_url in seen_urls:
                continue
            seen_urls.add(item_url)
            
            if len(title) < 5:
                title = unquote(item_url.split("/")[-1].replace("-", " ").capitalize())

            if any(w in title.lower() for w in GRADUATE_WORDS):
                continue

            pub_date = extract_date(title)
            clean_title = re.sub(r"^\d{1,2}[./-]\d{1,2}[./-](?:\d{4}|\d{2})\s*", "", title).strip()
            
            files = []
            lower_item_url = item_url.lower()
            if any(lower_item_url.endswith(ext) for ext in [".pdf", ".xlsx", ".xls"]):
                parsed = parse_file_exams(item_url)
                ext = lower_item_url.split(".")[-1].split("?")[0].lower()
                files.append({
                    "title": f"2026-2027 Güz Ara Sınav Programı ({ext.upper()})",
                    "url": item_url,
                    "ext": ext,
                    "exam_items": parsed
                })
            elif "/duyuru/detay/" in item_url:
                files = find_pdf_links_in_page(item_url, clean_title)
            
            # Aynı program hem PDF hem Excel olarak verildiyse Excel tercih edilir:
            # PDF tablolarında birden çok ders tek satıra yapışabiliyor.
            excel_files = [f for f in files if f.get("ext") in ("xlsx", "xls") and f.get("exam_items")]
            source_files = excel_files or files
            all_items = []
            for f in source_files:
                all_items.extend(it for it in f.get("exam_items", [])
                                 if len((it.get("code") or "").split()) <= 1)
            all_items = filter_for_department(all_items, dept_name)
            
            result["announcements"].append({
                "title": clean_title if clean_title else title,
                "date": pub_date or "Ekim 2026",
                "url": item_url,
                "files": files,
                "exam_items": all_items
            })
            
        if result["announcements"]:
            # Sayfa sırası güvenilir değil: tarihe göre en yeni program duyurusu başa gelir
            result["announcements"].sort(key=announcement_sort_key, reverse=True)
            result["has_announcement"] = True
            latest = result["announcements"][0]
            if not latest["exam_items"] and "kesinle" in latest["title"].lower():
                own, common_files = common_faculty_items(faculty_name, dept_name)
                if own:
                    latest["exam_items"] = own
                    latest["files"] = latest["files"] + common_files
            status, label = classify_status(result["announcements"])
            result["status"] = status
            result["status_label"] = label
            
    except Exception as e:
        print(f"[{dept_key}] Hata: {e}")
        
    return result

def main():
    print("=== BTÜ 2026-2027 Güz Ara Sınav Programları Taranıyor (V3) ===")
    all_results = {}
    
    total_depts = len(DEPARTMENTS)
    announced_count = 0
    finalized_count = 0
    draft_count = 0
    not_announced_count = 0
    
    for key, info in DEPARTMENTS.items():
        if key in SHARED_DEPARTMENTS and SHARED_DEPARTMENTS[key] in all_results:
            # Programı paylaşılan bölüm: kaynağın duyurularını kendi adıyla kullan
            src = all_results[SHARED_DEPARTMENTS[key]]
            data = dict(src, key=key, name=info.get("name", key),
                        faculty=FACULTY_MAP.get(key, src.get("faculty")))
        else:
            data = scrape_department_exams(key, info)
        all_results[key] = data
        
        st = data.get("status", "PENDING")
        if data.get("has_announcement"):
            announced_count += 1
            if st == "FINAL":
                finalized_count += 1
            elif st == "DRAFT":
                draft_count += 1
            ann = data["announcements"][0]
            items_cnt = len(ann.get("exam_items", []))
            print(f"[+] [{st}] {info['name']}: {data['status_label']} ({ann['title'][:40]}...) - {items_cnt} ders.")
        else:
            not_announced_count += 1
            print(f"[-] [PENDING] {info['name']}: Henüz duyurulmadı.")

    meta = {
        "term": "2026-2027 Güz Dönemi",
        "exam_type": "Ara Sınav (Vize)",
        "total_departments": total_depts,
        "announced_count": announced_count,
        "finalized_count": finalized_count,
        "draft_count": draft_count,
        "not_announced_count": not_announced_count,
        "last_updated": datetime.now().strftime("%d.%m.%Y %H:%M")
    }
    
    print("\n" + "="*50)
    print("GOSTERGE RAPORU:")
    print(f"[+] Kesinlesmis Programlar  : {finalized_count} Bolum")
    print(f"[+] Taslak Programlar       : {draft_count} Bolum")
    print(f"[-] Yayinlanmayan / Bekleyen: {not_announced_count} Bolum")
    print(f"[=] Toplam Taranan          : {total_depts} Bolum")
    print("="*50 + "\n")

    final_output = {
        "meta": meta,
        "departments": all_results
    }

    with open("exam_schedules.json", "w", encoding="utf-8") as f:
        json.dump(final_output, f, ensure_ascii=False, indent=2)
        
    js_content = f"// BTÜ 2026-2027 Güz Ara Sınav Programları - Güncellendi: {datetime.now().strftime('%d.%m.%Y %H:%M')}\n"
    js_content += f"window.EXAM_DATA = {json.dumps(final_output, ensure_ascii=False, indent=2)};\n"
    
    with open("exam_schedules.js", "w", encoding="utf-8") as f:
        f.write(js_content)
        
    # Tarayıcı eklentisi dizinine de kopyala
    if os.path.isdir("extension"):
        ext_target = os.path.join("extension", "exam_schedules.js")
        shutil.copyfile("exam_schedules.js", ext_target)
        print(f"[OK] {ext_target} güncellendi.")

    print(f"[OK] Tamamlandı: exam_schedules.json ve exam_schedules.js başarıyla kaydedildi.")

if __name__ == "__main__":
    main()
